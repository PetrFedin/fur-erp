# Единая книга «Расходы, сырьё и Cash Flow»: ОДИН массив данных (лист «Данные») для всех сводных + лист «Сверка» с цифрами приложения.
# Вход: прежняя книга (расходы и Cash Flow), raw_lots.json (scripts/export-raw-lots-json.ts), recon.json (scripts/export-recon-figures.ts).
# Сводные создаются XML (refreshOnLoad); после открытия в Excel: обновить и сохранить, чтобы данные легли внутрь книги.
#   python3 scripts/build-unified-workbook.py <old.xlsx> <raw_lots.json> <recon.json> <out.xlsx>
import datetime, json, re, sys, zipfile
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from xml.sax.saxutils import escape

old_path, raw_path, recon_path, out = sys.argv[1:5]
tmp = out + ".base.xlsx"
recon = json.load(open(recon_path)); raw = json.load(open(raw_path))
owb = openpyxl.load_workbook(old_path, data_only=True)

TEXT = ["Блок", "Источник", "Статус", "Год (по дате)", "Месяц", "Сезон (по дате)", "Дата", "Основание даты", "Тип меха", "Категория 1", "Категория 2", "Категория 3", "Артикул", "Код модели", "ТЗ", "Партия", "Год артикула", "Сезон артикула",
        "Основание сезона артикула", "Основание цены работы", "Направление", "Статья", "Месяц (CF)", "Вид оплаты", "Контрагент", "За что", "Документ", "Примечание", "Аукцион", "Лот", "Сорт (текст лота)", "Признаки меха", "Цвет"]
NUM = ["Изделий, шт.", "Мех (с расходами ТЗ)", "Подкладка", "Надбавка 4 %", "Работа", "Всего", "Мех без НДС, ₽", "Подкладка без НДС, ₽", "Надбавка 4 % без НДС, ₽", "Работа без НДС, ₽", "Всего без НДС, ₽",
       "Розничная цена, ₽/шт.", "Розничная цена без НДС, ₽/шт.", "Осталось по заказу, шт.", "Штуки заказа, шт.", "Цена работы на остаток, ₽/шт.", "Работа на остаток заказа, ₽", "Работа: произведено + запланировано, ₽",
       "Произведено, шт.", "Произведено с розн. ценой, шт.", "Всего с розн. ценой, ₽ с НДС", "Всего с розн. ценой, ₽ без НДС", "Розница по произведённому, ₽ с НДС", "Розница по произведённому, ₽ без НДС", "Работа по произведённому, ₽", "Маркап по строке, с НДС",
       "Сумма с НДС ₽", "Ставка НДС", "НДС ₽", "Сумма без НДС ₽", "Discount с НДС ₽", "Discount без НДС ₽", "Штуки",
       "Куплено, шт.", "Куплено, ₽ с НДС", "Использовано, шт.", "Использовано, ₽ с НДС", "Остаток, шт.", "Остаток, ₽ с НДС", "Куплено, ₽ без НДС", "Использовано, ₽ без НДС", "Остаток, ₽ без НДС",
       "Молоток, $/шт. (по лоту)", "Куплено с $, шт.", "Куплено, $", "Использовано с $, шт.", "Использовано, $", "Остаток с $, шт.", "Остаток, $"]
H = TEXT + NUM
IX = {n: i for i, n in enumerate(H)}
assert len(IX) == len(H)
BL_TZ, BL_ORD, BL_ARR, BL_CF, BL_RAW = "Расходы: ТЗ", "Расходы: заказ FW26", "Приход (поступление)", "Cash Flow", "Сырьё"

def blank_row(): return [None] * len(H)
def put(r, name, v): r[IX[name]] = v
def txt(v):
    if v is None or v == "": return "—"
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v)

rows = []
# --- расходы / заказ / приход
ws = owb["Данные для расходы"]
oh = [c.value for c in ws[1]]
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(oh, r))
    src = d["Источник"]
    if src == "Оплата (Cashflow)": continue
    blk = {"Расходы по ТЗ": BL_TZ, "Заказ FW26 (не произведено)": BL_ORD, "Поступление (выгрузка продаж)": BL_ARR}[src]
    row = blank_row(); put(row, "Блок", blk)
    for n in TEXT:
        if n in d and n != "Блок": put(row, n, txt(d[n]))
    for n in NUM:
        if n in d and d[n] is not None and not isinstance(d[n], str): put(row, n, d[n])
    for n in ("Год (по дате)", "Год артикула"): put(row, n, txt(d.get(n)))
    rows.append(row)
# ТЗ-ремонт (штуки в производство не входят, деньги входят в себестоимость, как на Главной)
for o in recon["repair"]:
    row = blank_row(); put(row, "Блок", BL_TZ)
    put(row, "Источник", "Расходы по ТЗ"); put(row, "Статус", "В работе" if o["st"] == "IN_PROGRESS" else "План")
    put(row, "Год (по дате)", str(o["y"])); put(row, "Месяц", "Без даты"); put(row, "Сезон (по дате)", "SS26"); put(row, "Основание даты", "ТЗ ремонт: даты нет")
    put(row, "Тип меха", o["fur"]); put(row, "Категория 1", "Ремонт"); put(row, "Категория 2", "Ремонт пальто"); put(row, "Категория 3", "—"); put(row, "Артикул", o["ml"]); put(row, "Код модели", "—")
    put(row, "ТЗ", o["number"]); put(row, "Партия", o["bl"]); put(row, "Год артикула", str(o["y"])); put(row, "Сезон артикула", "SS26"); put(row, "Основание сезона артикула", "ТЗ ремонт: штуки не считаются, деньги входят в себестоимость")
    put(row, "Изделий, шт.", 0); put(row, "Произведено, шт.", 0)
    mex, mexe = o["furIn"] + o["exp"], o["furEx"] + o["exp"]
    put(row, "Мех (с расходами ТЗ)", mex); put(row, "Подкладка", o["linIn"]); put(row, "Надбавка 4 %", o["c4In"]); put(row, "Работа", o["work"])
    put(row, "Всего", mex + o["linIn"] + o["c4In"] + o["work"])
    put(row, "Мех без НДС, ₽", mexe); put(row, "Подкладка без НДС, ₽", o["linEx"]); put(row, "Надбавка 4 % без НДС, ₽", o["c4Ex"]); put(row, "Работа без НДС, ₽", o["work"])
    put(row, "Всего без НДС, ₽", mexe + o["linEx"] + o["c4Ex"] + o["work"]); put(row, "Работа по произведённому, ₽", o["work"])
    rows.append(row)
# --- Cash Flow
cf = owb["Данные CF"]; ch = [c.value for c in cf[1]]
cf_rows = []
for r in cf.iter_rows(min_row=2, values_only=True):
    d = dict(zip(ch, r))
    row = blank_row(); put(row, "Блок", BL_CF); put(row, "Источник", "Cash Flow")
    put(row, "Дата", txt(d["Дата"])); put(row, "Год (по дате)", txt(d["Год"])); put(row, "Месяц (CF)", txt(d["Месяц"]))
    for n in ("Направление", "Статья", "Категория 1", "Категория 2", "Тип меха", "Вид оплаты", "Контрагент", "За что", "Документ", "Примечание"): put(row, n, txt(d[n]))
    for n in ("Сумма с НДС ₽", "Ставка НДС", "НДС ₽", "Сумма без НДС ₽", "Discount с НДС ₽", "Discount без НДС ₽", "Штуки"):
        if isinstance(d[n], (int, float)): put(row, n, d[n])
    cf_rows.append(row)
cf_rows.sort(key=lambda x: (x[IX["Месяц (CF)"]], x[IX["Направление"]]))
rows += cf_rows
# --- сырьё
for r in raw:
    row = blank_row(); put(row, "Блок", BL_RAW); put(row, "Источник", r["src"]); put(row, "Статус", "—")
    put(row, "Год (по дате)", str(r["year"])); put(row, "Тип меха", r["item"]); put(row, "Аукцион", txt(r["auction"])); put(row, "Лот", txt(r["lot"])); put(row, "Сорт (текст лота)", txt(r["grade"]))
    put(row, "Признаки меха", txt(r["attrs"])); put(row, "Цвет", txt(r["color"]))
    for n, k in (("Куплено, шт.", "recv"), ("Куплено, ₽ с НДС", "recvIn"), ("Использовано, шт.", "used"), ("Использовано, ₽ с НДС", "usedIn"), ("Остаток, шт.", "rem"), ("Остаток, ₽ с НДС", "remIn"),
                 ("Куплено, ₽ без НДС", "recvEx"), ("Использовано, ₽ без НДС", "usedEx"), ("Остаток, ₽ без НДС", "remEx"), ("Молоток, $/шт. (по лоту)", "usd"), ("Куплено с $, шт.", "usdQty"), ("Куплено, $", "usdSub"),
                 ("Использовано с $, шт.", "usdUQty"), ("Использовано, $", "usdUSub"), ("Остаток с $, шт.", "usdRQty"), ("Остаток, $", "usdRSub")):
        if r.get(k) is not None: put(row, n, r[k])
    rows.append(row)
# пустые текстовые -> «—»
for row in rows:
    for n in TEXT:
        if row[IX[n]] is None: row[IX[n]] = "—"
N = len(rows); LAST = N + 1
print("строк:", N, "по блокам:", {b: sum(1 for x in rows if x[0] == b) for b in (BL_TZ, BL_ORD, BL_ARR, BL_CF, BL_RAW)})

# ---------- книга
wb = openpyxl.Workbook()
SHEETS = ["Расходы с НДС", "Расходы без НДС", "Сырьё с НДС", "Сырьё без НДС", "Cash Flow", "Сверка", "Справка", "Данные"]
wb.active.title = SHEETS[0]
for s in SHEETS[1:]: wb.create_sheet(s)
dws = wb["Данные"]
dws.append(H)
for r in rows: dws.append(r)
for c in dws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="center")
dws.freeze_panes = "A2"; dws.auto_filter.ref = f"A1:{get_column_letter(len(H))}{LAST}"
for i, n in enumerate(H, 1):
    dws.column_dimensions[get_column_letter(i)].width = 14 if n in NUM else (30 if n in ("Основание даты", "Основание сезона артикула", "Основание цены работы", "Сорт (текст лота)", "Примечание", "За что", "Артикул") else 16)
    if n in NUM:
        for c in dws[get_column_letter(i)][1:]: c.number_format = "#,##0.00" if n in ("Молоток, $/шт. (по лоту)", "Куплено, $", "Использовано, $", "Остаток, $", "Маркап по строке, с НДС") else "#,##0.0" if n.endswith("шт.") or n == "Изделий, шт." else "#,##0"

def col(name): return get_column_letter(IX[name] + 1)
def rng(name): return f"Данные!${col(name)}$2:${col(name)}${LAST}"
def sumifs(name, *cond):
    s = f"SUMIFS({rng(name)}"
    for k, v in cond: s += f",{rng(k)},\"{v}\""
    return s + ")"

# ---------- Сверка
sw = wb["Сверка"]
sw.append(["Сверка книги с приложением (Главная / Оплаты / Cash Flow)"]); sw["A1"].font = Font(bold=True, size=13)
sw.append(["Колонка «Excel» считается формулами по листу «Данные». Колонка «Приложение» — цифры приложения на момент выгрузки (scripts/export-recon-figures.ts). Допуск: до 5 ₽ или 0,001 %."])
sw.append([])
sw.append(["Показатель", "Excel", "Приложение", "Разница", "Статус", "Комментарий"])
for c in sw[4]: c.font = Font(bold=True)
P = recon["prod"]; R = recon["raw"]; PD = recon["paid"]
items = [
    ("ТЗ: изделий, шт. (без ремонта)", sumifs("Изделий, шт.", ("Блок", BL_TZ)), P["qty"], "Главная → Производство, все сезоны"),
    ("ТЗ: мех с расходами ТЗ, ₽ с НДС", sumifs("Мех (с расходами ТЗ)", ("Блок", BL_TZ)), P["fur"], "включая ТЗ ремонт"),
    ("ТЗ: подкладка, ₽ с НДС", sumifs("Подкладка", ("Блок", BL_TZ)), P["lining"], ""),
    ("ТЗ: надбавка 4 %, ₽ с НДС", sumifs("Надбавка 4 %", ("Блок", BL_TZ)), P["c4"], ""),
    ("ТЗ: работа, ₽", sumifs("Работа", ("Блок", BL_TZ)), P["work"], "работа без НДС и с НДС одинакова"),
    ("ТЗ: полная себестоимость, ₽ с НДС", sumifs("Всего", ("Блок", BL_TZ)), P["fur"] + P["lining"] + P["c4"] + P["work"], "Главная → «Полная себестоимость · Всего»"),
    ("ТЗ: мех без НДС, ₽", sumifs("Мех без НДС, ₽", ("Блок", BL_TZ)), P["furEx"], ""),
    ("ТЗ: подкладка без НДС, ₽", sumifs("Подкладка без НДС, ₽", ("Блок", BL_TZ)), P["liningEx"], ""),
    ("ТЗ: надбавка 4 % без НДС, ₽", sumifs("Надбавка 4 % без НДС, ₽", ("Блок", BL_TZ)), P["c4Ex"], ""),
    ("Сырьё: куплено, шт.", sumifs("Куплено, шт.", ("Блок", BL_RAW)), R["recvQty"], "Главная → Сырьё → Мех, все года; с «заказано, не получено»"),
    ("Сырьё: куплено, ₽ с НДС", sumifs("Куплено, ₽ с НДС", ("Блок", BL_RAW)), R["recvValue"], "разница — округление по лотам"),
    ("Сырьё: использовано, шт.", sumifs("Использовано, шт.", ("Блок", BL_RAW)), R["usedQty"], ""),
    ("Сырьё: использовано, ₽ с НДС", sumifs("Использовано, ₽ с НДС", ("Блок", BL_RAW)), R["usedValue"], "без надбавки 4 %"),
    ("Сырьё: остаток, шт.", sumifs("Остаток, шт.", ("Блок", BL_RAW)), R["remQty"], ""),
    ("Сырьё: остаток, ₽ с НДС", sumifs("Остаток, ₽ с НДС", ("Блок", BL_RAW)), R["remValue"], ""),
    ("Сырьё: куплено, ₽ без НДС", sumifs("Куплено, ₽ без НДС", ("Блок", BL_RAW)), R["recvValueEx"], ""),
    ("Сырьё: использовано, ₽ без НДС", sumifs("Использовано, ₽ без НДС", ("Блок", BL_RAW)), R["usedValueEx"], ""),
    ("Сырьё: остаток, ₽ без НДС", sumifs("Остаток, ₽ без НДС", ("Блок", BL_RAW)), R["remValueEx"], ""),
    ("Оплачено: мех, ₽ с НДС", "(" + sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Мех: аукционы")) + "+" + sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Мех: закупка у поставщика")) + "+" + sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Шкуры подрядчика в счёте за пошив")) + ")", PD["fur"], "Главная → Оплачено всего → Мех"),
    ("Оплачено: подкладка, ₽ с НДС", sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Подкладка")), PD["lining"], ""),
    ("Оплачено: услуги по УПД, ₽ с НДС", sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Услуги по УПД")), PD["services"], ""),
    ("Оплачено: работа подрядчику, ₽ с НДС", sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата"), ("Статья", "Работа подрядчику")), PD["work"], "включая аванс на будущие ТЗ"),
    ("Оплачено всего, ₽ с НДС", sumifs("Сумма с НДС ₽", ("Блок", BL_CF), ("Направление", "Трата")), PD["fur"] + PD["lining"] + PD["services"] + PD["work"], "Главная → Оплачено всего"),
]
for i, (name, f, app, cm) in enumerate(items, 5):
    sw.append([name, "=" + f, app, f"=B{i}-C{i}", f'=IF(ABS(D{i})<=MAX(5,0.00001*ABS(C{i})),"совпадает","РАСХОЖДЕНИЕ")', cm])
    sw[f"B{i}"].number_format = sw[f"C{i}"].number_format = sw[f"D{i}"].number_format = "#,##0"
sw.append([]); sw.append(["Не сверяется с приложением напрямую: средние (ср. работа, ср. розн. цена, маркап), разрезы по сезонам FW26, «Получено» в Cash Flow — независимого числа в приложении для них нет."])
for k, w in zip("ABCDEF", (44, 18, 18, 14, 16, 60)): sw.column_dimensions[k].width = w
# справка
rw = wb["Справка"]; rw.column_dimensions["A"].width = 34; rw.column_dimensions["B"].width = 120
rw.append(["Что", "Пояснение"]); rw["A1"].font = rw["B1"].font = Font(bold=True)
for a, b in [("Один массив данных", "Лист «Данные» — единая таблица. Колонка «Блок» разделяет строки: «Расходы: ТЗ», «Расходы: заказ FW26», «Приход (поступление)», «Cash Flow», «Сырьё». Все сводные берут эту таблицу и фильтруют блок."),
             ("Расходы с НДС / без НДС", "Блок «Расходы: ТЗ» + «Расходы: заказ FW26». Сезон артикула по умолчанию FW26. Средние: работа, себестоимость, розн. цена (округлена до 100 ₽) и маркап считаются по произведённому товару."),
             ("Сырьё с НДС / без НДС", "Блок «Сырьё»: лоты меха + «Заказано, не получено» (считается купленным, как на Главной). Ср. молоток в $ считается только по лотам, где цена в долларах известна."),
             ("Cash Flow", "Блок «Cash Flow»: по месяцам, статьи, суммы с НДС / НДС / без НДС. «Трата» и «Получено» — отдельные направления."),
             ("Сверка", "Excel против приложения: себестоимость ТЗ, сырьё, оплаты. Статус «совпадает», если разница не больше 5 ₽ или 0,001 %."),
             ("ТЗ ремонт", "ТЗ ремонта (П-2026-0183) входит в блок «Расходы: ТЗ» с нулевыми штуками: деньги в себестоимости, штуки не считаются, как на Главной."),
             ("Обновление", "Данные выгружаются скриптами (export-raw-lots-json.ts, export-recon-figures.ts, build-unified-workbook.py). После правок на листе «Данные»: Данные → Обновить всё.")]: rw.append([a, b])
for r in rw.iter_rows(min_row=2): r[1].alignment = Alignment(wrap_text=True, vertical="top")
wb.save(tmp)

# ---------- сводные (XML)
def uniq(name):
    seen = []
    i = IX[name]
    for r in rows:
        v = r[i]
        if v not in seen: seen.append(v)
    return seen
AXIS = ["Блок", "Источник", "Статус", "Год (по дате)", "Месяц", "Тип меха", "Категория 1", "Категория 2", "Сезон артикула", "Год артикула", "Направление", "Статья", "Месяц (CF)", "Признаки меха", "Цвет", "Аукцион", "Лот"]
SV = {a: uniq(a) for a in AXIS}
CALC = [
    ("Ср. работа на шт. (произведено), ₽", "'Работа по произведённому, ₽'/'Произведено, шт.'"),
    ("Ср. розн. цена (произведено), ₽ с НДС", "ROUND('Розница по произведённому, ₽ с НДС'/'Произведено с розн. ценой, шт.',-2)"),
    ("Ср. розн. цена (произведено), ₽ без НДС", "ROUND('Розница по произведённому, ₽ без НДС'/'Произведено с розн. ценой, шт.',-2)"),
    ("Ср. себестоимость на шт. с НДС (произведено), ₽", "'Всего'/'Произведено, шт.'"),
    ("Ср. себестоимость на шт. без НДС (произведено), ₽", "'Всего без НДС, ₽'/'Произведено, шт.'"),
    ("Маркап (произведено), с НДС", "('Розница по произведённому, ₽ с НДС'/'Произведено с розн. ценой, шт.')/('Всего'/'Произведено, шт.')"),
    ("Маркап (произведено), без НДС", "('Розница по произведённому, ₽ без НДС'/'Произведено с розн. ценой, шт.')/('Всего без НДС, ₽'/'Произведено, шт.')"),
    ("Ср. цена закупки, ₽/шт. с НДС", "'Куплено, ₽ с НДС'/'Куплено, шт.'"), ("Ср. цена использования, ₽/шт. с НДС", "'Использовано, ₽ с НДС'/'Использовано, шт.'"), ("Ср. цена остатка, ₽/шт. с НДС", "'Остаток, ₽ с НДС'/'Остаток, шт.'"),
    ("Ср. цена закупки, ₽/шт. без НДС", "'Куплено, ₽ без НДС'/'Куплено, шт.'"), ("Ср. цена использования, ₽/шт. без НДС", "'Использовано, ₽ без НДС'/'Использовано, шт.'"), ("Ср. цена остатка, ₽/шт. без НДС", "'Остаток, ₽ без НДС'/'Остаток, шт.'"),
    ("Ср. молоток закупки, $", "'Куплено, $'/'Куплено с $, шт.'"), ("Ср. молоток использования, $", "'Использовано, $'/'Использовано с $, шт.'"), ("Ср. молоток остатка, $", "'Остаток, $'/'Остаток с $, шт.'"),
]
NF = len(H); NC = len(CALC); TOT = NF + NC
CI = {n: NF + i for i, (n, _) in enumerate(CALC)}
def fld(n): return IX[n] if n in IX else CI[n]
def cache_field(i):
    n = H[i]; nm = escape(n)
    if n in SV: return f'<cacheField name="{nm}" numFmtId="0"><sharedItems count="{len(SV[n])}">' + "".join(f'<s v="{escape(v)}"/>' for v in SV[n]) + '</sharedItems></cacheField>'
    if n in NUM: return f'<cacheField name="{nm}" numFmtId="0"><sharedItems containsString="0" containsBlank="1" containsNumber="1"/></cacheField>'
    return f'<cacheField name="{nm}" numFmtId="0"><sharedItems/></cacheField>'
cache = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" refreshOnLoad="1" refreshedVersion="8" minRefreshableVersion="3" createdVersion="6" recordCount="{N}">'
         f'<cacheSource type="worksheet"><worksheetSource ref="A1:{get_column_letter(NF)}{LAST}" sheet="Данные"/></cacheSource><cacheFields count="{TOT}">'
         + "".join(cache_field(i) for i in range(NF)) + "".join(f'<cacheField name="{escape(n_)}" numFmtId="0" formula="{escape(f_, {chr(34): "&quot;"})}" databaseField="0"/>' for n_, f_ in CALC) + '</cacheFields></pivotCacheDefinition>')

def pivot(name, rows_f, cols_f, pages, datas, loc, caption, row_data_pos=False, col_data=False):
    """pages: [(поле, [выбранные значения] | None)] ; datas: [(подпись, поле, numFmt)]"""
    axis = {}
    for f in rows_f: axis[f] = "axisRow"
    for f in cols_f: axis[f] = "axisCol"
    for f, _ in pages: axis[f] = "axisPage"
    dsets = {fld(f) for _, f, _ in datas}
    pf = []
    for i in range(TOT):
        n = H[i] if i < NF else CALC[i - NF][0]
        if n in axis:
            vals = SV[n]; sel = dict(pages).get(n)
            multi = n in dict(pages) and sel is not None and len(sel) > 1
            its = "".join(f'<item h="1" x="{k}"/>' if (n in dict(pages) and sel is not None and len(sel) > 1 and v not in sel) else f'<item x="{k}"/>' for k, v in enumerate(vals))
            pf.append(f'<pivotField axis="{axis[n]}"' + (' multipleItemSelectionAllowed="1"' if multi else '') + f' showAll="0"><items count="{len(vals)+1}">{its}<item t="default"/></items></pivotField>')
        elif i in dsets: pf.append('<pivotField dataField="1" showAll="0"/>' if i < NF else '<pivotField dataField="1" dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        elif i >= NF: pf.append('<pivotField dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        else: pf.append('<pivotField showAll="0"/>')
    rf = "".join(f'<field x="{IX[f]}"/>' for f in rows_f) + ('<field x="-2"/>' if row_data_pos else "")
    cfs = ("".join(f'<field x="{IX[f]}"/>' for f in cols_f) + ('<field x="-2"/>' if col_data else '')) if cols_f else '<field x="-2"/>'
    pgs = ""
    for f, sel in pages:
        item = ""
        if sel is not None and len(sel) == 1: item = f' item="{SV[f].index(sel[0])}"'
        pgs += f'<pageField fld="{IX[f]}"{item} hier="-1"/>'
    dfs = "".join(f'<dataField name="{escape(c)}" fld="{fld(f)}" baseField="0" baseItem="0" numFmtId="{nf}"/>' for c, f, nf in datas)
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="{escape(name)}" cacheId="1" showError="1" errorCaption="—" applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="Значения" updatedVersion="8" minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0" outline="1" outlineData="1" multipleFieldFilters="0" rowHeaderCaption="{escape(caption)}">'
            f'<location ref="{loc}" firstHeaderRow="1" firstDataRow="{3 if col_data else 2}" firstDataCol="1" rowPageCount="{len(pages)}" colPageCount="1"/><pivotFields count="{TOT}">{"".join(pf)}</pivotFields><rowFields count="{len(rows_f) + (1 if row_data_pos else 0)}">{rf}</rowFields><colFields count="{(len(cols_f) + (1 if col_data else 0)) if cols_f else 1}">{cfs}</colFields>'
            f'<pageFields count="{len(pages)}">{pgs}</pageFields><dataFields count="{len(datas)}">{dfs}</dataFields><pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/></pivotTableDefinition>')

EXP_PAGES = [("Блок", [BL_TZ, BL_ORD]), ("Год артикула", None), ("Месяц", None), ("Сезон артикула", ["FW26"]), ("Статус", None)]
def exp_datas(vat):
    mex, lin, c4, wk, tot = (("Мех (с расходами ТЗ)", "Подкладка", "Надбавка 4 %", "Работа", "Всего") if vat else ("Мех без НДС, ₽", "Подкладка без НДС, ₽", "Надбавка 4 % без НДС, ₽", "Работа без НДС, ₽", "Всего без НДС, ₽"))
    s = "с НДС" if vat else "без НДС"
    return [("Работа на остаток заказа, ₽", "Работа на остаток заказа, ₽", 3), ("Штуки остатка заказа, шт.", "Осталось по заказу, шт.", 3), ("Штуки заказа, шт.", "Штуки заказа, шт.", 3), ("Работа по всему заказу, ₽", "Работа: произведено + запланировано, ₽", 3),
            ("Произведено, шт.", "Произведено, шт.", 3), ("Мех, ₽", mex, 3), ("Подкладка, ₽", lin, 3), ("Надбавка 4 %, ₽", c4, 3), ("Работа по ТЗ, ₽", wk, 3), ("Себестоимость произведённого, ₽", tot, 3),
            ("Ср. работа на шт., ₽ (произведено)", "Ср. работа на шт. (произведено), ₽", 3), ("Ср. себестоимость на шт., ₽ (произведено)", "Ср. себестоимость на шт. с НДС (произведено), ₽" if vat else "Ср. себестоимость на шт. без НДС (произведено), ₽", 3),
            ("Ср. розн. цена, ₽ (произведено)", f"Ср. розн. цена (произведено), ₽ {s}", 3), ("Маркап (произведено)", f"Маркап (произведено), {s}", 2)]
def raw_datas(vat):
    s = "с НДС" if vat else "без НДС"
    def blk(label, gen, q, rub, u):
        return [(f"{label}, шт.", q, 3), (f"{label}, ₽", rub, 3), (f"Ср. молоток {gen}, $ (по лотам с $)", f"Ср. молоток {gen}, $", 3), (f"Ср. цена {gen}, ₽/шт.", f"Ср. цена {gen}, ₽/шт. {s}", 3)]
    return (blk("Куплено", "закупки", "Куплено, шт.", "Куплено, ₽ с НДС" if vat else "Куплено, ₽ без НДС", 0) + blk("Использовано", "использования", "Использовано, шт.", "Использовано, ₽ с НДС" if vat else "Использовано, ₽ без НДС", 0) + blk("Остаток", "остатка", "Остаток, шт.", "Остаток, ₽ с НДС" if vat else "Остаток, ₽ без НДС", 0))
RAW_ROWS = ["Тип меха", "Признаки меха", "Цвет", "Источник", "Год (по дате)", "Аукцион", "Лот"]
specs = [
    ("Расходы с НДС", ["Категория 1", "Тип меха", "Категория 2"], [], EXP_PAGES, exp_datas(True), "A9:O60", "Категория / мех / категория 2", False),
    ("Расходы без НДС", ["Категория 1", "Тип меха", "Категория 2"], [], EXP_PAGES, exp_datas(False), "A9:O60", "Категория / мех / категория 2", False),
    ("Сырьё с НДС", RAW_ROWS, [], [("Блок", [BL_RAW])], raw_datas(True), "A5:N60", "Мех / признаки / цвет / источник / год / аукцион / лот", False),
    ("Сырьё без НДС", RAW_ROWS, [], [("Блок", [BL_RAW])], raw_datas(False), "A5:N60", "Мех / признаки / цвет / источник / год / аукцион / лот", False),
    ("Cash Flow по месяцам", ["Направление", "Статья"], ["Месяц (CF)"], [("Блок", [BL_CF]), ("Год (по дате)", None)], [("Сумма с НДС ₽", "Сумма с НДС ₽", 3), ("НДС ₽ (в составе)", "НДС ₽", 3), ("Сумма без НДС ₽", "Сумма без НДС ₽", 3)], "A6:V40", "Направление / статья", False, True),
]
# у «Блок» с одним значением: если выбрано одно — item=
import os
_sel = [int(x) for x in os.environ.get('PIVOTS', '').split(',') if x] or list(range(len(specs)))
tables = [pivot(*specs[i]) for i in _sel]
SHEETN = [i + 1 for i in _sel]
zi = zipfile.ZipFile(tmp); zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
CT = "".join(f'<Override PartName="/xl/pivotTables/pivotTable{k}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>' for k in range(1, len(tables) + 1))
for nm in zi.namelist():
    b = zi.read(nm)
    if nm == "[Content_Types].xml":
        b = b.decode().replace("</Types>", '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>' + CT + "</Types>").encode()
    elif nm == "xl/workbook.xml":
        b = b.decode().replace('</workbook>', '<pivotCaches><pivotCache xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" cacheId="1" r:id="rIdPC1"/></pivotCaches></workbook>').encode()
    elif nm == "xl/_rels/workbook.xml.rels":
        b = b.decode().replace("</Relationships>", '<Relationship Id="rIdPC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode()
    zo.writestr(nm, b)
zo.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
rel = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'
for k, t in enumerate(tables, 1):
    zo.writestr(f"xl/pivotTables/pivotTable{k}.xml", t); zo.writestr(f"xl/pivotTables/_rels/pivotTable{k}.xml.rels", rel)
    zo.writestr(f"xl/worksheets/_rels/sheet{SHEETN[k-1]}.xml.rels", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" Target="../pivotTables/pivotTable{k}.xml"/></Relationships>')
zo.close()
print("ok", N)
