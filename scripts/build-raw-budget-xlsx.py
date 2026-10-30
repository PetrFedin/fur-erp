#!/usr/bin/env python3
"""Excel «Сырьё: куплено, потрачено, остаток; остаток FW26; план и закупка FW27» — 7 листов, сводные и расчёт формулами.

  python3 -I scripts/build-raw-budget-xlsx.py <lots.json> <data.json> <out.xlsx>

lots.json — scripts/export-raw-lots-json.ts (лоты как на Главной, включая «Заказано, не получено» и шкуры подрядчика);
data.json — scripts/export-raw-budget-data.ts (прогноз по артикулам остатка FW26, склад, цены, бюджет FW27, Budget PCS).
Жёлтые ячейки — ввод: меняете — пересчитываются сводные и закупка.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

lots = json.load(open(sys.argv[1], encoding="utf-8"))
data = json.load(open(sys.argv[2], encoding="utf-8"))
OUT = sys.argv[3]

NAME = {"Свакара": "Каракульча"}  # как в бюджете и Budget PCS
FUR_ORDER = ["Соболь", "Норка", "Каракульча", "Лиса", "Рысь"]
fur_of = lambda f: NAME.get(f, f)

# ---------- стиль ----------
INK = "1F2933"
F_HEAD = PatternFill("solid", fgColor="E8ECF1")
F_FUR = PatternFill("solid", fgColor="D9E2EC")
F_GRP = PatternFill("solid", fgColor="EEF2F6")
F_TOT = PatternFill("solid", fgColor="BCCCDC")
F_IN = PatternFill("solid", fgColor="FFF3BF")
thin = Side(style="thin", color="C7CED6")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
BOLD = Font(bold=True, color=INK)
TITLE = Font(bold=True, size=14, color=INK)
NOTE = Font(italic=True, size=9, color="52606D")
INT, RUB, USD, PCT, DEC = '#,##0;-#,##0;"–"', '#,##0" ₽";-#,##0" ₽";"–"', '#,##0.00" $";-#,##0.00" $";"–"', '0%', '#,##0.0;-#,##0.0;"–"'
WRAP = Alignment(wrap_text=True, vertical="top")


def head(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c)
        x.font, x.fill, x.border, x.alignment = BOLD, F_HEAD, B, Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def title(ws, text, notes):
    ws["A1"] = text
    ws["A1"].font = TITLE
    for i, n in enumerate(notes, 2):
        ws.cell(i, 1, n).font = NOTE


def put(ws, r, c, v, fmt=None, fill=None, bold=False):
    x = ws.cell(r, c, v)
    x.border = B
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if bold:
        x.font = BOLD
    return x


def inp(ws, r, c, v, fmt=None):
    x = put(ws, r, c, v, fmt, F_IN)
    x.font = Font(bold=True, color="7C5E10")
    return x


wb = Workbook()

# =====================================================================
# Лист 1. Лоты
# =====================================================================
ws1 = wb.active
ws1.title = "1 Лоты"
title(ws1, "Всё купленное сырьё по лотам", [
    f"Данные на {data['asOf']} (МСК). Как на Главной: лоты склада + «Заказано, не получено» + шкуры подрядчика. Суммы в ₽ — с НДС.",
    "Цена $ — цена молотка за шкуру на аукционе (там, где она есть в документах); средневзвешенная по количеству. Подкладка и прочие материалы в лист не входят.",
])
C1 = ["Тип меха", "Год", "Аукцион / поставщик", "Источник", "Лот", "Сорт / описание", "Цвет",
      "Куплено, шк.", "Ср. цена $", "Ср. цена ₽", "Сумма куплено, ₽",
      "Потрачено, шк.", "Потрачено, ₽", "Остаток, шк.", "Остаток, ₽",
      "Шк. с ценой $", "Сумма $", "Потрачено шк. с $", "Потрачено $", "Остаток шк. с $", "Остаток $"]
HR1 = 5
head(ws1, HR1, C1, [12, 7, 20, 20, 16, 34, 14, 10, 11, 12, 15, 10, 15, 10, 15, 10, 13, 10, 13, 10, 13])
rows = sorted(lots, key=lambda r: (FUR_ORDER.index(fur_of(r["item"])) if fur_of(r["item"]) in FUR_ORDER else 99, r["item"], r["year"], r["auction"], r["lot"]))
r = HR1 + 1
for x in rows:
    put(ws1, r, 1, fur_of(x["item"]))
    put(ws1, r, 2, int(x["year"]) if str(x["year"]).isdigit() else x["year"])
    put(ws1, r, 3, x["auction"])
    put(ws1, r, 4, x["src"])
    put(ws1, r, 5, x["lot"])
    put(ws1, r, 6, x["grade"])
    put(ws1, r, 7, x["color"])
    put(ws1, r, 8, x["recv"], INT)
    put(ws1, r, 9, f"=IFERROR(Q{r}/P{r},0)", USD)
    put(ws1, r, 10, f"=IFERROR(K{r}/H{r},0)", RUB)
    put(ws1, r, 11, x["recvIn"], RUB)
    put(ws1, r, 12, x["used"], INT)
    put(ws1, r, 13, x["usedIn"], RUB)
    put(ws1, r, 14, x["rem"], INT)
    put(ws1, r, 15, x["remIn"], RUB)
    put(ws1, r, 16, x["usdQty"], INT)
    put(ws1, r, 17, x["usdSub"], USD)
    put(ws1, r, 18, x["usdUQty"], INT)
    put(ws1, r, 19, x["usdUSub"], USD)
    put(ws1, r, 20, x["usdRQty"], INT)
    put(ws1, r, 21, x["usdRSub"], USD)
    r += 1
LAST1 = r - 1
put(ws1, r, 1, "ИТОГО", fill=F_TOT, bold=True)
for c in range(2, 8):
    put(ws1, r, c, None, fill=F_TOT)
for c in [8, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21]:
    put(ws1, r, c, f"=SUBTOTAL(9,{L(c)}{HR1+1}:{L(c)}{LAST1})", ws1.cell(LAST1, c).number_format, F_TOT, True)
put(ws1, r, 9, f"=IFERROR(Q{r}/P{r},0)", USD, F_TOT, True)
put(ws1, r, 10, f"=IFERROR(K{r}/H{r},0)", RUB, F_TOT, True)
ws1.auto_filter.ref = f"A{HR1}:{L(len(C1))}{LAST1}"
ws1.freeze_panes = f"H{HR1+1}"
LOT = lambda col: f"'1 Лоты'!${col}${HR1+1}:${col}${LAST1}"

# =====================================================================
# Лист 2. Сводная: тип меха → год
# =====================================================================
ws2 = wb.create_sheet("2 Сводная сырьё")
title(ws2, "Сводная: куплено, потрачено, остаток по типам меха и годам", [
    "Формулы от листа «1 Лоты». Ср. цена ₽ = сумма ₽ ÷ шкуры; ср. цена $ — только по шкурам с ценой молотка.",
])
C2 = ["Тип меха", "Год",
      "Куплено, шк.", "Ср. $", "Ср. ₽", "Сумма, ₽",
      "Потрачено, шк.", "Ср. $", "Ср. ₽", "Сумма, ₽",
      "Остаток, шк.", "Ср. $", "Ср. ₽", "Сумма, ₽"]
HR2 = 4
head(ws2, HR2, C2, [14, 8, 11, 10, 11, 15, 11, 10, 11, 15, 11, 10, 11, 15])
for c0, t in [(3, "КУПЛЕНО"), (7, "ПОТРАЧЕНО"), (11, "ОСТАТОК")]:
    ws2.merge_cells(start_row=HR2 - 1, start_column=c0, end_row=HR2 - 1, end_column=c0 + 3)
    x = ws2.cell(HR2 - 1, c0, t)
    x.font, x.alignment, x.fill = BOLD, Alignment(horizontal="center"), F_GRP
furs_all = []
for x in rows:
    f = fur_of(x["item"])
    if f not in furs_all:
        furs_all.append(f)
years = sorted({int(x["year"]) for x in rows if str(x["year"]).isdigit()})
# пары: (шк., сумма ₽, шк. с $, сумма $) для блоков куплено / потрачено / остаток
BLK = [("H", "K", "P", "Q"), ("L", "M", "R", "S"), ("N", "O", "T", "U")]


def block_formulas(ws, r, crit):
    for i, (q, rub, uq, us) in enumerate(BLK):
        c = 3 + i * 4
        put(ws, r, c, f"=SUMIFS({LOT(q)},{crit})", INT)
        put(ws, r, c + 1, f"=IFERROR(SUMIFS({LOT(us)},{crit})/SUMIFS({LOT(uq)},{crit}),0)", USD)
        put(ws, r, c + 3, f"=SUMIFS({LOT(rub)},{crit})", RUB)
        put(ws, r, c + 2, f"=IFERROR({L(c+3)}{r}/{L(c)}{r},0)", RUB)


r = HR2 + 1
fur_rows = []
for f in furs_all:
    put(ws2, r, 1, f, fill=F_FUR, bold=True)
    put(ws2, r, 2, "все годы", fill=F_FUR, bold=True)
    block_formulas(ws2, r, f'{LOT("A")},$A{r}')
    for c in range(3, 15):
        ws2.cell(r, c).fill, ws2.cell(r, c).font = F_FUR, BOLD
    fur_rows.append(r)
    fr = r
    r += 1
    for y in years:
        if not any(fur_of(x["item"]) == f and str(x["year"]) == str(y) for x in rows):
            continue
        put(ws2, r, 1, f)
        ws2.cell(r, 1).font = Font(color="9AA5B1")
        put(ws2, r, 2, y)
        block_formulas(ws2, r, f'{LOT("A")},$A{r},{LOT("B")},$B{r}')
        r += 1
put(ws2, r, 1, "ВСЕГО", fill=F_TOT, bold=True)
put(ws2, r, 2, "", fill=F_TOT)
for i, (q, rub, uq, us) in enumerate(BLK):
    c = 3 + i * 4
    put(ws2, r, c, "=" + "+".join(f"{L(c)}{x}" for x in fur_rows), INT, F_TOT, True)
    put(ws2, r, c + 3, "=" + "+".join(f"{L(c+3)}{x}" for x in fur_rows), RUB, F_TOT, True)
    put(ws2, r, c + 2, f"=IFERROR({L(c+3)}{r}/{L(c)}{r},0)", RUB, F_TOT, True)
    put(ws2, r, c + 1, f"=IFERROR(SUM({LOT(us)})/SUM({LOT(uq)}),0)", USD, F_TOT, True)
ws2.freeze_panes = f"C{HR2+1}"
STOCK_ROW = {f: fr for f, fr in zip(furs_all, fur_rows)}  # строка типа меха: остаток шк. — столбец K, сумма — N

# =====================================================================
# Лист 3. Артикулы остатка FW26 в производстве
# =====================================================================
ws3 = wb.create_sheet("3 Артикулы в производстве")
title(ws3, "Артикулы заказа FW26, которые ещё в производстве: сколько шкур и денег ещё уйдёт", [
    f"Прогноз «{data.get('forecastName') or '—'}»: осталось сшить = объём ТЗ модели − получено. Шкур на изделие и цены — из ТЗ артикула или средние по категории (столбец «Источник»).",
    "«Уже выдано» — шкуры и подкладка, списанные со склада по ТЗ этой модели, которые ещё в работе: они уже не на складе и второй раз не считаются.",
    "Жёлтые ячейки можно менять (например, вписать расход шкур там, где «нет данных»). «Норка + Соболь» делится на норку и соболь на листе 4.",
])
C3 = ["Артикул", "Одежда / аксессуары", "Категория", "Тип меха", "Осталось сшить, шт.", "Шкур на изделие", "Уже выдано по ТЗ в работе, шк.", "Ещё уйдёт шкур",
      "Цена шкуры, ₽", "Сырьё, ₽", "Подкладка, м/изд.", "Уже выдано, м", "Цена м, ₽", "Подкладка, ₽", "Работа, ₽/изд.", "Работа, ₽", "Прочее, ₽", "Итого ещё затрат, ₽", "Источник"]
HR3 = 5
head(ws3, HR3, C3, [26, 12, 18, 16, 10, 10, 12, 10, 12, 14, 10, 10, 10, 12, 12, 14, 11, 15, 30])
fc = sorted(data["forecast"], key=lambda l: (l["fur"], l["cls"], l["acode"]))
r = HR3 + 1
for l in fc:
    f = l["fur"] if " + " in l["fur"] else fur_of(l["fur"])
    put(ws3, r, 1, l["acode"])
    put(ws3, r, 2, l["cls"])
    put(ws3, r, 3, l["cat2"])
    put(ws3, r, 4, f)
    put(ws3, r, 5, l["left"], INT)
    inp(ws3, r, 6, l["hp"] if l["hp"] else None, DEC)
    put(ws3, r, 7, round(l.get("issued") or 0, 2), DEC)
    put(ws3, r, 8, f"=MAX(0,E{r}*F{r}-G{r})", DEC)
    hp_price = l["hidePrice"] or data["price"].get(f) or None
    inp(ws3, r, 9, round(hp_price) if hp_price else None, RUB)
    put(ws3, r, 10, f"=H{r}*I{r}", RUB)
    inp(ws3, r, 11, l["mp"] or None, DEC)
    put(ws3, r, 12, round(l.get("issuedM") or 0, 2), DEC)
    inp(ws3, r, 13, round(l["meterPrice"]) if l["meterPrice"] else None, RUB)
    put(ws3, r, 14, f"=MAX(0,E{r}*K{r}-L{r})*M{r}", RUB)
    inp(ws3, r, 15, round(l["wp"]) if l["wp"] else None, RUB)
    put(ws3, r, 16, f"=E{r}*O{r}", RUB)
    inp(ws3, r, 17, l["extra"] or None, RUB)
    put(ws3, r, 18, f"=J{r}+N{r}+P{r}+Q{r}", RUB)
    src = l["source"] + (" · работа ориентировочно" if l.get("wpEst") else "") + ("" if l["hp"] else " · НЕТ РАСХОДА ШКУР — впишите")
    put(ws3, r, 19, src)
    r += 1
LAST3 = r - 1
put(ws3, r, 1, "ИТОГО", fill=F_TOT, bold=True)
for c in range(2, 20):
    put(ws3, r, c, None, fill=F_TOT)
for c in [5, 7, 8, 10, 12, 14, 16, 17, 18]:
    put(ws3, r, c, f"=SUBTOTAL(9,{L(c)}{HR3+1}:{L(c)}{LAST3})", ws3.cell(LAST3, c).number_format, F_TOT, True)
ws3.auto_filter.ref = f"A{HR3}:{L(len(C3))}{LAST3}"
ws3.freeze_panes = f"B{HR3+1}"
ART = lambda col: f"'3 Артикулы в производстве'!${col}${HR3+1}:${col}${LAST3}"

# =====================================================================
# Лист 4. Сводная: что уйдёт на остаток FW26 и что останется
# =====================================================================
ws4 = wb.create_sheet("4 Остаток FW26 сводная")
title(ws4, "Сколько шкур и денег ещё уйдёт на оставшиеся артикулы FW26 и сколько сырья останется", [
    "На складе — остаток с листа 2 (включая заказанное, но не полученное). «Уйдёт» — без шкур, уже выданных по ТЗ в работе (они уже списаны со склада). «Норка + Соболь» делится по фактической доле норки в ТЗ (жёлтые ячейки).",
])
ws4["A4"], ws4["A5"] = "Доля норки в «Норка + Соболь», одежда", "Доля норки в «Норка + Соболь», аксессуары"
for a in ("A4", "A5"):
    ws4[a].font = BOLD
inp(ws4, 4, 3, round(data["mix"]["RTW"], 4), PCT)
inp(ws4, 5, 3, round(data["mix"]["Acc"], 4), PCT)
MIX_RTW, MIX_ACC = "'4 Остаток FW26 сводная'!$C$4", "'4 Остаток FW26 сводная'!$C$5"
C4 = ["Тип меха", "На складе сейчас, шк.", "Сумма склада, ₽", "Уйдёт на остаток FW26, шк.", "из них одежда", "из них аксессуары",
      "Сырьё на остаток FW26, ₽", "Подкладка, работа, прочее, ₽", "Всего затрат на остаток FW26, ₽", "Останется после производства, шк.", "Останется, ₽ (по ср. цене склада)", "Не хватает, шк."]
HR4 = 7
head(ws4, HR4, C4, [32, 13, 15, 13, 11, 11, 15, 15, 16, 14, 16, 11])
r = HR4 + 1
rows4 = []


def need_formula(fur, col_block=None):
    """Шкуры на остаток FW26 по типу меха (с долей норки/соболя из «Норка + Соболь»); col_block — 'Одежда'/'Аксессуары' или None."""
    blk = f',{ART("B")},"{col_block}"' if col_block else ""
    own = f'SUMIFS({ART("H")},{ART("D")},"{fur}"{blk})'
    if fur not in ("Норка", "Соболь"):
        return own
    part = lambda b, mix: f'SUMIFS({ART("H")},{ART("D")},"Норка + Соболь",{ART("B")},"{b}")*{mix if fur == "Норка" else "(1-" + mix + ")"}'
    if col_block == "Одежда":
        return f"{own}+{part('Одежда', MIX_RTW)}"
    if col_block == "Аксессуары":
        return f"{own}+{part('Аксессуары', MIX_ACC)}"
    return f"{own}+{part('Одежда', MIX_RTW)}+{part('Аксессуары', MIX_ACC)}"


def money_formula(fur, col):
    own = f'SUMIFS({ART(col)},{ART("D")},"{fur}")'
    if fur not in ("Норка", "Соболь"):
        return own
    # деньги смеси делим по той же доле шкур
    sh = lambda b, mix: f'SUMIFS({ART(col)},{ART("D")},"Норка + Соболь",{ART("B")},"{b}")*{mix if fur == "Норка" else "(1-" + mix + ")"}'
    return f"{own}+{sh('Одежда', MIX_RTW)}+{sh('Аксессуары', MIX_ACC)}"


for f in furs_all:
    sr = STOCK_ROW[f]
    put(ws4, r, 1, f, bold=f in FUR_ORDER)
    put(ws4, r, 2, f"='2 Сводная сырьё'!K{sr}", INT)
    put(ws4, r, 3, f"='2 Сводная сырьё'!N{sr}", RUB)
    put(ws4, r, 4, f"={need_formula(f)}", INT)
    put(ws4, r, 5, f"={need_formula(f, 'Одежда')}", INT)
    put(ws4, r, 6, f"={need_formula(f, 'Аксессуары')}", INT)
    put(ws4, r, 7, f"={money_formula(f, 'J')}", RUB)
    put(ws4, r, 8, f"={money_formula(f, 'N')}+{money_formula(f, 'P')}+{money_formula(f, 'Q')}", RUB)
    put(ws4, r, 9, f"=G{r}+H{r}", RUB)
    put(ws4, r, 10, f"=MAX(0,B{r}-D{r})", INT)
    put(ws4, r, 11, f"=IFERROR(J{r}*C{r}/B{r},0)", RUB)
    put(ws4, r, 12, f"=MAX(0,D{r}-B{r})", INT)
    rows4.append(r)
    r += 1
put(ws4, r, 1, "ВСЕГО", fill=F_TOT, bold=True)
for c, fmt in zip(range(2, 13), [INT, RUB, INT, INT, INT, RUB, RUB, RUB, INT, RUB, INT]):
    put(ws4, r, c, f"=SUM({L(c)}{HR4+1}:{L(c)}{r-1})", fmt, F_TOT, True)
TOT4 = r
r += 2
ws4.cell(r, 1, "Проверка: шкур на остаток FW26 по листу 3 (без деления смеси)").font = NOTE
put(ws4, r, 4, f"=SUM({ART('H')})", INT)
ws4.cell(r + 1, 1, "Проверка: затраты на остаток FW26 по листу 3").font = NOTE
put(ws4, r + 1, 9, f"=SUM({ART('R')})", RUB)
ws4.freeze_panes = f"B{HR4+1}"
LEFT_ROW = {f: x for f, x in zip(furs_all, rows4)}

# =====================================================================
# Лист 5. План 2027: что хотим произвести
# =====================================================================
ws5 = wb.create_sheet("5 План 2027")
title(ws5, "План производства 2027 (бюджет FW27): категории, тип меха, длина, количество и расход меха", [
    "Одежда — как в бюджете (RTW от 19.10). Аксессуары — новые количества по файлу MVST W FURS ACC от 27.10 (905 шт.).",
    "Шкур на изделие — из бюджета FW27 (статистика ТЗ по категории и длине; курсивом в бюджете — примерная оценка по аналогам, см. столбец «Оценка»).",
])
C5 = ["Тип меха", "Одежда / аксессуары", "Категория", "Длина", "Кол-во, шт.", "Шкур на изделие", "Шкур всего", "Подкладка, м/изд.", "Подкладка всего, м", "Оценка (нет данных системы)"]
HR5 = 5
head(ws5, HR5, C5, [14, 13, 24, 10, 10, 11, 11, 11, 12, 60])
FO = ["Соболь", "Норка", "Норка+Соболь", "Лиса", "Каракульча", "Рысь"]
bud = sorted(data["budget"], key=lambda b: (FO.index(b["fur"]) if b["fur"] in FO else 99, 0 if b["group"] == "Одежда" else 1, b["cat"], ["Short", "Regular", "Long", ""].index(b["len"]) if b["len"] in ["Short", "Regular", "Long", ""] else 9))
r = HR5 + 1
for b in bud:
    put(ws5, r, 1, b["fur"])
    put(ws5, r, 2, b["group"])
    put(ws5, r, 3, b["cat"])
    put(ws5, r, 4, b["len"] or "—")
    inp(ws5, r, 5, b["pcs"], INT)
    inp(ws5, r, 6, round(b["hides"], 2), DEC)
    put(ws5, r, 7, f"=E{r}*F{r}", INT)
    put(ws5, r, 8, round(b["meters"] or 0, 2), DEC)
    put(ws5, r, 9, f"=E{r}*H{r}", DEC)
    put(ws5, r, 10, b["est"] or "")
    if b["est"]:
        ws5.cell(r, 6).font = Font(italic=True, bold=True, color="7C5E10")
    r += 1
LAST5 = r - 1
put(ws5, r, 1, "ИТОГО", fill=F_TOT, bold=True)
for c in range(2, 11):
    put(ws5, r, c, None, fill=F_TOT)
for c in [5, 7, 9]:
    put(ws5, r, c, f"=SUBTOTAL(9,{L(c)}{HR5+1}:{L(c)}{LAST5})", ws5.cell(LAST5, c).number_format, F_TOT, True)
put(ws5, r, 6, f"=IFERROR(G{r}/E{r},0)", DEC, F_TOT, True)
r += 2
ws5.cell(r, 1, "Итого по типам меха и блокам").font = BOLD
r += 1
head(ws5, r, ["Тип меха", "Одежда / аксессуары", "", "", "Кол-во, шт.", "Шкур на изделие", "Шкур всего"])
P5 = lambda col: f"${col}${HR5+1}:${col}${LAST5}"
r += 1
for f in FO:
    for g in ["Одежда", "Аксессуары"]:
        put(ws5, r, 1, f)
        put(ws5, r, 2, g)
        put(ws5, r, 5, f'=SUMIFS({P5("E")},{P5("A")},A{r},{P5("B")},B{r})', INT)
        put(ws5, r, 7, f'=SUMIFS({P5("G")},{P5("A")},A{r},{P5("B")},B{r})', INT)
        put(ws5, r, 6, f"=IFERROR(G{r}/E{r},0)", DEC)
        r += 1
ws5.auto_filter.ref = f"A{HR5}:{L(len(C5))}{LAST5}"
ws5.freeze_panes = f"E{HR5+1}"

# =====================================================================
# Лист 6. Budget PCS 2027: деньги на сырьё
# =====================================================================
ws6 = wb.create_sheet("6 Budget PCS 2027")
title(ws6, "MVST · Budget PCS 2027: сколько сырья и денег на весь сезон (без учёта остатков)", [
    "Количество и шкур на изделие — с листа 5. Ср. цена шкуры — из бюджета FW27 по категории (средневзвешенная по ТЗ); рост цены задаётся сверху.",
])
ws6["A4"] = "Рост цены шкуры, %"
ws6["A4"].font = BOLD
inp(ws6, 4, 3, 0.10, PCT)
UP = "'6 Budget PCS 2027'!$C$4"
C6 = ["Уровень", "Тип меха", "Категория 1", "Категория 2", "Длина", "Кол-во, шт.", "Ср. шкур на изд.", "Шкур всего", "Ср. цена шкуры, ₽", "Цена с ростом, ₽", "На сырьё, ₽"]
HR6 = 6
head(ws6, HR6, C6, [11, 14, 13, 24, 10, 10, 11, 11, 13, 13, 16])
# листья
r = HR6 + 1
leaf_first = r
for i, b in enumerate(bud):
    src = HR5 + 1 + i  # строка листа 5
    put(ws6, r, 1, "строка")
    put(ws6, r, 2, b["fur"])
    put(ws6, r, 3, b["group"])
    put(ws6, r, 4, b["cat"])
    put(ws6, r, 5, b["len"] or "—")
    put(ws6, r, 6, f"='5 План 2027'!E{src}", INT)
    put(ws6, r, 7, f"='5 План 2027'!F{src}", DEC)
    put(ws6, r, 8, f"=F{r}*G{r}", INT)
    inp(ws6, r, 9, round(b["hidePrice"] or 0), RUB)
    put(ws6, r, 10, f"=I{r}*(1+{UP})", RUB)
    put(ws6, r, 11, f"=H{r}*J{r}", RUB)
    r += 1
leaf_last = r - 1
P6 = lambda col: f"'6 Budget PCS 2027'!${col}${leaf_first}:${col}${leaf_last}"
r += 1
ws6.cell(r, 1, "Сводная как в отчёте MVST: TOTAL → тип меха → категория 1 → категория 2").font = BOLD
r += 1
head(ws6, r, C6)
r += 1


def agg_row(r, level, crit, labels, fill=None):
    put(ws6, r, 1, level, fill=fill, bold=bool(fill))
    for c, v in zip([2, 3, 4, 5], labels):
        put(ws6, r, c, v, fill=fill, bold=bool(fill))
    sf = lambda col: f"SUMIFS({P6(col)},{crit})" if crit else f"SUM({P6(col)})"
    put(ws6, r, 6, f"={sf('F')}", INT, fill, bool(fill))
    put(ws6, r, 8, f"={sf('H')}", INT, fill, bool(fill))
    put(ws6, r, 11, f"={sf('K')}", RUB, fill, bool(fill))
    put(ws6, r, 7, f"=IFERROR(H{r}/F{r},0)", DEC, fill, bool(fill))
    put(ws6, r, 10, f"=IFERROR(K{r}/H{r},0)", RUB, fill, bool(fill))
    put(ws6, r, 9, f"=IFERROR(J{r}/(1+{UP}),0)", RUB, fill, bool(fill))


agg_row(r, "TOTAL", None, ["Все", "", "", ""], F_TOT)
TOT6 = r
r += 1
for f in FO:
    if not any(b["fur"] == f for b in bud):
        continue
    agg_row(r, "тип меха", f'{P6("B")},"{f}"', [f, "", "", ""], F_FUR)
    r += 1
    for g in ["Одежда", "Аксессуары"]:
        cats = []
        for b in bud:
            if b["fur"] == f and b["group"] == g and b["cat"] not in cats:
                cats.append(b["cat"])
        if not cats:
            continue
        agg_row(r, "категория 1", f'{P6("B")},"{f}",{P6("C")},"{g}"', [f, g, "", ""], F_GRP)
        r += 1
        for c in cats:
            agg_row(r, "категория 2", f'{P6("B")},"{f}",{P6("C")},"{g}",{P6("D")},"{c}"', [f, g, c, ""])
            r += 1
ws6.freeze_panes = f"F{HR6+1}"

# =====================================================================
# Лист 7. Расчёт закупки с остатками и запасом
# =====================================================================
ws7 = wb.create_sheet("7 Расчёт закупки")
title(ws7, "Сколько сырья докупить на план 2027 с учётом остатка после FW26, качества остатка и запаса", [
    "Порядок: остаток после производства FW26 (лист 4) × годная доля = можно использовать; годное делится: аксессуарам — их доля, одежде — остальное (лучшая часть).",
    "Нужно из новой покупки = шкур на план (лист 6) − можно использовать. Закупить = нужно ÷ (1 − запас): покупаем с запасом, в ноль не покупаем. Рысь — по потребности блоков.",
    "Сверху — значения по умолчанию; в таблице «Параметры по типам меха» их можно переопределить для каждого типа (жёлтые ячейки).",
])
G = {}
params = [("Годная доля остатка, %", 0.70, "g_use"), ("Годная доля остатка, рысь, %", 0.85, "g_use_lynx"), ("Из годного остатка — в аксессуары, %", 0.75, "g_acc"),
          ("Запас при закупке (остаток от купленного), %", 0.30, "g_res"), ("Запас при закупке, рысь, %", 0.15, "g_res_lynx")]
for i, (t, v, k) in enumerate(params):
    rr = 6 + i
    ws7.cell(rr, 1, t).font = BOLD
    inp(ws7, rr, 4, v, PCT)
    G[k] = f"$D${rr}"
ws7.cell(11, 1, "Рост цены шкуры, %").font = BOLD
put(ws7, 11, 4, f"={UP}", PCT)
ws7.cell(11, 5, "← задаётся на листе 6").font = NOTE
G["up"] = "$D$11"
# параметры по типам меха
r = 13
ws7.cell(r, 1, "Параметры по типам меха (можно менять каждый)").font = BOLD
r += 1
head(ws7, r, ["Тип меха", "Годная доля остатка", "В аксессуары из годного", "Запас при закупке", "Ср. цена шкуры, ₽ (база)"])
r += 1
PR = {}
for f in FUR_ORDER:
    put(ws7, r, 1, f, bold=True)
    lynx = f == "Рысь"
    inp(ws7, r, 2, f"={G['g_use_lynx'] if lynx else G['g_use']}", PCT)
    if lynx:
        put(ws7, r, 3, "по потребности")
    else:
        inp(ws7, r, 3, f"={G['g_acc']}", PCT)
    inp(ws7, r, 4, f"={G['g_res_lynx'] if lynx else G['g_res']}", PCT)
    inp(ws7, r, 5, round(data["price"].get(f, 0)), RUB)
    PR[f] = r
    r += 1
r += 1
C7 = ["Тип меха", "Блок", "Изделий на план, шт.", "Шкур на план (TTL)", "Шкур на изделие", "Остаток после FW26, шк.", "Годная доля", "Можно использовать, шк.",
      "Нужно из новой покупки, шк.", "Запас", "Закупить с запасом, шк.", "Останется после 2027 (запас), шк.", "Ср. цена шкуры, ₽", "Цена с ростом, ₽", "На закупку, ₽"]
HR7 = r
head(ws7, HR7, C7, [34, 12, 12, 12, 10, 13, 9, 13, 13, 8, 13, 13, 12, 12, 16])
r += 1


def ttl6(fur, blk):
    base = f'SUMIFS({P6("H")},{P6("B")},"{fur}",{P6("C")},"{blk}")'
    if fur in ("Норка", "Соболь"):
        mix = MIX_RTW if blk == "Одежда" else MIX_ACC
        share = mix if fur == "Норка" else f"(1-{mix})"
        base += f'+SUMIFS({P6("H")},{P6("B")},"Норка+Соболь",{P6("C")},"{blk}")*{share}'
    return base


def pcs6(fur, blk):
    return f'SUMIFS({P6("F")},{P6("B")},"{fur}",{P6("C")},"{blk}")'


tot_rows = []
for f in FUR_ORDER:
    p = PR[f]
    lr = LEFT_ROW.get(f)
    left_total = f"'4 Остаток FW26 сводная'!J{lr}" if lr else "0"
    short_total = f"'4 Остаток FW26 сводная'!L{lr}" if lr else "0"
    first = r
    for blk in ["Одежда", "Аксессуары"]:
        put(ws7, r, 1, f, bold=True)
        put(ws7, r, 2, blk)
        put(ws7, r, 3, f"={pcs6(f, blk)}", INT)
        put(ws7, r, 4, f"={ttl6(f, blk)}", INT)
        put(ws7, r, 5, f"=IFERROR(D{r}/C{r},0)", DEC)
        if f == "Рысь":
            share = f"IFERROR(D{r}/({ttl6(f, 'Одежда')}+{ttl6(f, 'Аксессуары')}),0.5)"
        else:
            share = f"$C${p}" if blk == "Аксессуары" else f"(1-$C${p})"
        put(ws7, r, 6, f"={left_total}*{share}", INT)
        put(ws7, r, 7, f"=$B${p}", PCT)
        put(ws7, r, 8, f"=F{r}*G{r}", INT)
        put(ws7, r, 9, f"=MAX(0,D{r}-H{r})+{short_total}*{share}", INT)
        put(ws7, r, 10, f"=$D${p}", PCT)
        put(ws7, r, 11, f"=IFERROR(I{r}/(1-J{r}),0)", INT)
        put(ws7, r, 12, f"=K{r}-I{r}+MAX(0,H{r}-D{r})", INT)
        put(ws7, r, 13, f"=$E${p}", RUB)
        put(ws7, r, 14, f"=M{r}*(1+{G['up']})", RUB)
        put(ws7, r, 15, f"=K{r}*N{r}", RUB)
        r += 1
    put(ws7, r, 1, f"{f} — итого", fill=F_FUR, bold=True)
    put(ws7, r, 2, "", fill=F_FUR)
    for c, fmt in [(3, INT), (4, INT), (6, INT), (8, INT), (9, INT), (11, INT), (12, INT), (15, RUB)]:
        put(ws7, r, c, f"=SUM({L(c)}{first}:{L(c)}{r-1})", fmt, F_FUR, True)
    put(ws7, r, 5, f"=IFERROR(D{r}/C{r},0)", DEC, F_FUR, True)
    for c in (7, 10, 13):
        put(ws7, r, c, None, fill=F_FUR)
    put(ws7, r, 14, f"=IFERROR(O{r}/K{r},0)", RUB, F_FUR, True)
    tot_rows.append(r)
    r += 1
put(ws7, r, 1, "ВСЕГО (норка и соболь включают свою долю «Норка + Соболь»)", fill=F_TOT, bold=True)
put(ws7, r, 2, "", fill=F_TOT)
for c, fmt in [(3, INT), (4, INT), (6, INT), (8, INT), (9, INT), (11, INT), (12, INT), (15, RUB)]:
    put(ws7, r, c, "=" + "+".join(f"{L(c)}{x}" for x in tot_rows), fmt, F_TOT, True)
put(ws7, r, 14, f"=IFERROR(O{r}/K{r},0)", RUB, F_TOT, True)
for c in (5, 7, 10, 13):
    put(ws7, r, c, None, fill=F_TOT)
put(ws7, r, 5, f"=IFERROR(D{r}/C{r},0)", DEC, F_TOT, True)
TOT7 = r
r += 1
ws7.cell(r, 1, "«Изделий на план» у норки и соболя не включает изделия «Норка + Соболь» (их шкуры — включены). Изделий «Норка + Соболь»:").font = NOTE
put(ws7, r, 3, f'=SUMIFS({P6("F")},{P6("B")},"Норка+Соболь")', INT)
r += 2
ws7.cell(r, 1, "Кратко").font = TITLE
r += 1
for t, fml, fmt in [
    ("Сырьё на складе сейчас, шк. (лист 4)", f"='4 Остаток FW26 сводная'!B{TOT4}", INT),
    ("Уйдёт на остаток заказа FW26, шк.", f"='4 Остаток FW26 сводная'!D{TOT4}", INT),
    ("Останется после FW26, шк.", f"='4 Остаток FW26 сводная'!J{TOT4}", INT),
    ("Нужно шкур на план 2027 (TTL)", f"=D{TOT7}", INT),
    ("Из остатка можно использовать, шк.", f"=H{TOT7}", INT),
    ("Закупить с запасом, шк.", f"=K{TOT7}", INT),
    ("На закупку сырья, ₽ (с ростом цены, с НДС)", f"=O{TOT7}", RUB),
    ("Для сравнения: сырьё на весь план без остатков, ₽ (лист 6)", f"='6 Budget PCS 2027'!K{TOT6}", RUB),
]:
    ws7.cell(r, 1, t)
    put(ws7, r, 4, fml, fmt, bold=True)
    r += 1
ws7.freeze_panes = f"C{HR7+1}"

for ws in wb.worksheets:
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("сохранено:", OUT)
