# Книга «Каракуль: сводные»: лист «Данные» (лоты) + две настоящие сводные (с НДС / без НДС) + «Справка».
# Вход: JSON из scripts/export-karakul-report.ts (DUMP_JSON). Сводные создаются XML-ом (refreshOnLoad), после открытия в Excel их нужно один раз обновить и сохранить.
#   python3 scripts/build-karakul-pivot.py <in.json> <out.xlsx>
import json, re, shutil, sys, zipfile
import openpyxl
from openpyxl.styles import Font, Alignment
from xml.sax.saxutils import escape

rows = json.load(open(sys.argv[1])); out = sys.argv[2]
ORDER = {"Свакара": 0, "Каракуль": 1, "Бродтейл": 2}
rows.sort(key=lambda r: (ORDER.get(r["g"], 9), 0 if "аукционе" in r["src"] else 1, r["year"], r["date"])); tmp = out + ".base.xlsx"
H = ["Вид", "Лот", "Дата", "Год", "Аукцион", "Сорт (текст лота)", "Куплено, шт.", "Цена закупки, ₽/шт. с НДС", "Куплено, ₽ с НДС", "Использовано, шт.", "Использовано, ₽ с НДС", "Остаток, шт.", "Остаток, ₽ с НДС", "Прочее списание, шт.", "Куплено, ₽ без НДС", "Использовано, ₽ без НДС", "Остаток, ₽ без НДС", "Источник"]
keys = ["g", "lot", "date", "year", "auction", "grade", "recv", "unit", "recvIn", "used", "usedIn", "rem", "remIn", "other", "recvEx", "usedEx", "remEx", "src"]
wb = openpyxl.Workbook()
wb.active.title = "Сводная с НДС"; wb.create_sheet("Сводная без НДС"); ref = wb.create_sheet("Справка"); ws = wb.create_sheet("Данные")
ws.append(H)
ws.cell(row=1, column=19, value="Основание вида")
for r in rows: ws.append([r[k] for k in keys] + [r.get("basis", "")])
ws.column_dimensions["R"].width = 24; ws.column_dimensions["S"].width = 46
for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([12, 14, 12, 8, 22, 46, 12, 14, 16, 14, 16, 12, 16, 12, 16, 16, 16, 24], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
for col in "GJLN": 
    for c in ws[col][1:]: c.number_format = "#,##0.0"
for col in "HIKMOPQ":
    for c in ws[col][1:]: c.number_format = "#,##0"
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:R{len(rows)+1}"
ref.column_dimensions["A"].width = 30; ref.column_dimensions["B"].width = 110
ref.append(["Термин", "Что значит"]); ref["A1"].font = ref["B1"].font = Font(bold=True)
for a, b in [("Каракуль", "Шкурки ягнят каракульской породы овец. Общее название меха: плотный курчавый «завиток». Здесь — лоты, где в сорте сказано просто «каракуль», без слов swakara и broadtail."),
 ("Свакара (Swakara)", "Торговая марка намибийского каракуля (Swakara — Karakul Pelts of Namibia). По типу тот же каракуль, с упором на происхождение и качество; шкурки с чёткими завитками, часто крашеные."),
 ("Бродтейл (Broadtail, каракульча)", "Мех самых молодых каракулят: ягнят забивают в первые дни жизни или недоношенными. Шкурка тонкая, лёгкая, шёлковистая, с гладким муаровым (волнистым) рисунком без завитков; дороже и нежнее каракуля. В справочнике системы пока сведён в один вид «Свакара»."),
 ("Куплено / Использовано / Остаток", "Куплено — получено по лотам закупки. Использовано — списано в ТЗ. Остаток — лежит на складе. Прочее списание = получено − использовано − остаток."),
 ("Как отнесены лоты", "По тексту сорта лота: «SWAKARA/свакара» → Свакара; «каракульча/broadtail» → Бродтейл. Лоты 728001, 728002, 728011, 728026 (994 шт., SAGA 06-2025) в УПД названы «шкурки каракуля», но в расчёте закупки это KARAKUL BROADTAIL GREY — отнесены к Бродтейлу. Каракуль без пометок swakara/broadtail — отдельно (сейчас лотов нет)."),
 ("Как пользоваться", "Сводные на первых двух листах: «Вид» (с итогом по виду) → «Куплено на аукционе» / «Куплено у поставщика» → годы; фильтр «Аукцион». Данные — на листе «Данные»; после правок: Данные → Обновить всё.")]: ref.append([a, b])
for r in ref.iter_rows(min_row=2): r[1].alignment = Alignment(wrap_text=True, vertical="top")
wb.save(tmp)

n = len(rows); last = n + 1
def uniq(i): 
    seen = []; 
    for r in rows:
        v = str(r[keys[i]])
        if v not in seen: seen.append(v)
    return seen
sv = {i: uniq(i) for i in (0, 3, 4, 5, 17)}
num = {6: "General", 7: "General"}
def cache_field(i, name):
    if i in sv: return f'<cacheField name="{escape(name)}" numFmtId="0"><sharedItems count="{len(sv[i])}">' + "".join(f'<s v="{escape(v)}"/>' for v in sv[i]) + '</sharedItems></cacheField>'
    if i in (1, 2): return f'<cacheField name="{escape(name)}" numFmtId="0"><sharedItems/></cacheField>'
    return f'<cacheField name="{escape(name)}" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"/></cacheField>'
calc = [("Ср. цена закупки, ₽/шт. с НДС", "'Куплено, ₽ с НДС'/'Куплено, шт.'"), ("Ср. цена закупки, ₽/шт. без НДС", "'Куплено, ₽ без НДС'/'Куплено, шт.'")]
fields = "".join(cache_field(i, H[i]) for i in range(18)) + "".join(f'<cacheField name="{escape(n_)}" numFmtId="0" formula="{escape(f_)}" databaseField="0"/>' for n_, f_ in calc)
cache = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" refreshOnLoad="1" refreshedVersion="8" minRefreshableVersion="3" createdVersion="6" recordCount="{n}"><cacheSource type="worksheet"><worksheetSource ref="A1:R{last}" sheet="Данные"/></cacheSource><cacheFields count="20">{fields}</cacheFields></pivotCacheDefinition>'
def pfield(i):
    if i in (0, 3, 17): return f'<pivotField axis="axisRow" showAll="0"><items count="{len(sv[i])+1}">' + "".join(f'<item x="{k}"/>' for k in range(len(sv[i]))) + '<item t="default"/></items></pivotField>'
    if i in (4,): return f'<pivotField axis="axisPage" showAll="0"><items count="{len(sv[i])+1}">' + "".join(f'<item x="{k}"/>' for k in range(len(sv[i]))) + '<item t="default"/></items></pivotField>'
    return None
def table(name, data):  # data: [(name, fld)]
    pf = []
    dset = {f for _, f in data}
    for i in range(20):
        x = pfield(i)
        if x: pf.append(x)
        elif i in dset: pf.append('<pivotField dataField="1" showAll="0"/>' if i < 18 else '<pivotField dataField="1" dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        elif i >= 18: pf.append('<pivotField dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        else: pf.append('<pivotField showAll="0"/>')
    dfs = "".join(f'<dataField name="{escape(nm)}" fld="{f}" baseField="0" baseItem="0" numFmtId="{4 if f >= 18 else 3 if f not in (6, 9, 11) else 4}"/>' for nm, f in data)
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="{name}" cacheId="1" applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="Значения" updatedVersion="8" minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0" outline="1" outlineData="1" multipleFieldFilters="0" rowHeaderCaption="Вид / источник / год"><location ref="A6:I12" firstHeaderRow="0" firstDataRow="1" firstDataCol="1" rowPageCount="1" colPageCount="1"/><pivotFields count="20">{"".join(pf)}</pivotFields><rowFields count="3"><field x="0"/><field x="17"/><field x="3"/></rowFields><colFields count="1"><field x="-2"/></colFields><pageFields count="1"><pageField fld="4" hier="-1"/></pageFields><dataFields count="{len(data)}">{dfs}</dataFields><pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/></pivotTableDefinition>'
t1 = table("Каракуль с НДС", [("Куплено, шт.", 6), ("Куплено, ₽", 8), ("Ср. цена закупки, ₽/шт.", 18), ("Использовано, шт.", 9), ("Использовано, ₽", 10), ("Остаток, шт.", 11), ("Остаток, ₽", 12), ("Прочее списание, шт.", 13)])
t2 = table("Каракуль без НДС", [("Куплено, шт.", 6), ("Куплено, ₽", 14), ("Ср. цена закупки, ₽/шт.", 19), ("Использовано, шт.", 9), ("Использовано, ₽", 15), ("Остаток, шт.", 11), ("Остаток, ₽", 16), ("Прочее списание, шт.", 13)])
zi = zipfile.ZipFile(tmp); zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
names = zi.namelist(); sheets = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', zi.read("xl/workbook.xml").decode())
for nm in names:
    b = zi.read(nm)
    if nm == "[Content_Types].xml":
        t = b.decode(); add = '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/><Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/><Override PartName="/xl/pivotTables/pivotTable2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        b = t.replace("</Types>", add + "</Types>").encode()
    elif nm == "xl/workbook.xml":
        t = b.decode(); b = t.replace('</workbook>', '<pivotCaches><pivotCache xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" cacheId="1" r:id="rIdPC1"/></pivotCaches></workbook>').encode()
    elif nm == "xl/_rels/workbook.xml.rels":
        t = b.decode(); b = t.replace("</Relationships>", '<Relationship Id="rIdPC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode()
    zo.writestr(nm, b)
zo.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
zo.writestr("xl/pivotTables/pivotTable1.xml", t1); zo.writestr("xl/pivotTables/pivotTable2.xml", t2)
rel = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'
zo.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels", rel); zo.writestr("xl/pivotTables/_rels/pivotTable2.xml.rels", rel)
for idx, pt in ((1, 1), (2, 2)):
    p = f"xl/worksheets/_rels/sheet{idx}.xml.rels"
    zo.writestr(p, f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" Target="../pivotTables/pivotTable{pt}.xml"/></Relationships>')
zo.close()
print("ok", n)
