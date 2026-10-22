# Книга «Сырьё: сводные»: лист «Данные» (лоты меха) + две сводные (с НДС / без НДС) + «Справка».
# Вход: JSON из scripts/export-raw-lots-json.ts. Сводные создаются XML (refreshOnLoad), после открытия в Excel их нужно обновить и сохранить.
#   python3 scripts/build-raw-pivot.py <in.json> <out.xlsx>
import json, re, sys, zipfile
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from xml.sax.saxutils import escape

rows = json.load(open(sys.argv[1])); out = sys.argv[2]; tmp = out + ".base.xlsx"
rows.sort(key=lambda r: (r["item"], 0 if "аукционе" in r["src"] else 1, r["year"], r["auction"], r["lot"]))
H = ["Вид меха", "Источник", "Год", "Аукцион", "Лот", "Сорт (текст лота)", "Куплено, шт.", "Куплено, ₽ с НДС", "Использовано, шт.", "Использовано, ₽ с НДС", "Остаток, шт.", "Остаток, ₽ с НДС",
     "Куплено, ₽ без НДС", "Использовано, ₽ без НДС", "Остаток, ₽ без НДС", "Молоток, $/шт. (по лоту)", "Куплено с $, шт.", "Куплено, $", "Использовано с $, шт.", "Использовано, $", "Остаток с $, шт.", "Остаток, $"]
K = ["item", "src", "year", "auction", "lot", "grade", "recv", "recvIn", "used", "usedIn", "rem", "remIn", "recvEx", "usedEx", "remEx", "usd", "usdQty", "usdSub", "usdUQty", "usdUSub", "usdRQty", "usdRSub"]
NF = len(H)
wb = openpyxl.Workbook()
wb.active.title = "Сырьё с НДС"; wb.create_sheet("Сырьё без НДС"); ref = wb.create_sheet("Справка"); ws = wb.create_sheet("Данные")
ws.append(H)
for r in rows: ws.append([r[k] for k in K])
for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([14, 22, 7, 22, 14, 40, 11, 14, 12, 14, 11, 14, 14, 14, 14, 12, 12, 12, 12, 12, 12, 12], 1): ws.column_dimensions[get_column_letter(i)].width = w
for col in "GIKQSU":
    for c in ws[col][1:]: c.number_format = "#,##0.0"
for col in "HJLMNO":
    for c in ws[col][1:]: c.number_format = "#,##0"
for col in "PRTV":
    for c in ws[col][1:]: c.number_format = "#,##0.00"
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(NF)}{len(rows)+1}"
ref.column_dimensions["A"].width = 34; ref.column_dimensions["B"].width = 110
ref.append(["Показатель", "Как считается"]); ref["A1"].font = ref["B1"].font = Font(bold=True)
for a, b in [("Куплено / Использовано / Остаток", "Шкуры по лотам закупки (принятые на склад). Использовано — списано в ТЗ; остаток — на складе. Заказано, но не получено, сюда не входит."),
 ("Ср. цена закупки / использования / остатка", "Сумма ₽ ÷ количество шкур (средневзвешенная по штукам), с НДС или без НДС в зависимости от листа."),
 ("Молоток, $", "Цена молотка за шкуру в долларах по документам лота. Для лотов без цены в $ она не определена и в расчёты $ не входит."),
 ("Ср. молоток $ (куплено / использовано / остаток)", "Σ $ ÷ Σ шт. только по шкурам из лотов, где цена в $ известна. Колонки «Куплено с $, шт.» и т. д. показывают покрытие."),
 ("Куплено на аукционе / у поставщика", "Аукцион — закупка на пушном аукционе (SAGA, Kopenhagen, Fur Harvesters и т. д.); остальное и шкуры подрядчика — у поставщика."),
 ("Как смотреть", "Строки: вид меха → источник → аукцион → лот. Фильтр «Год». После изменений на листе «Данные»: Данные → Обновить всё.")]: ref.append([a, b])
for r in ref.iter_rows(min_row=2): r[1].alignment = Alignment(wrap_text=True, vertical="top")
wb.save(tmp)

n = len(rows); last = n + 1
def uniq(i):
    seen = []
    for r in rows:
        v = str(r[K[i]])
        if v not in seen: seen.append(v)
    return seen
AX_ROW = [0, 1, 3, 4]; AX_PAGE = [2]
sv = {i: uniq(i) for i in AX_ROW + AX_PAGE}
def cache_field(i):
    nm = escape(H[i])
    if i in sv: return f'<cacheField name="{nm}" numFmtId="0"><sharedItems count="{len(sv[i])}">' + "".join(f'<s v="{escape(v)}"/>' for v in sv[i]) + '</sharedItems></cacheField>'
    if i == 5: return f'<cacheField name="{nm}" numFmtId="0"><sharedItems/></cacheField>'
    return f'<cacheField name="{nm}" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"/></cacheField>' if i != 15 else f'<cacheField name="{nm}" numFmtId="0"><sharedItems containsString="0" containsBlank="1" containsNumber="1"/></cacheField>'
# вычисляемые поля: индексы NF..
CALC = [
    ("Ср. цена закупки, ₽/шт. с НДС", "'Куплено, ₽ с НДС'/'Куплено, шт.'"),
    ("Ср. цена использования, ₽/шт. с НДС", "'Использовано, ₽ с НДС'/'Использовано, шт.'"),
    ("Ср. цена остатка, ₽/шт. с НДС", "'Остаток, ₽ с НДС'/'Остаток, шт.'"),
    ("Ср. цена закупки, ₽/шт. без НДС", "'Куплено, ₽ без НДС'/'Куплено, шт.'"),
    ("Ср. цена использования, ₽/шт. без НДС", "'Использовано, ₽ без НДС'/'Использовано, шт.'"),
    ("Ср. цена остатка, ₽/шт. без НДС", "'Остаток, ₽ без НДС'/'Остаток, шт.'"),
    ("Ср. молоток закупки, $", "'Куплено, $'/'Куплено с $, шт.'"),
    ("Ср. молоток использования, $", "'Использовано, $'/'Использовано с $, шт.'"),
    ("Ср. молоток остатка, $", "'Остаток, $'/'Остаток с $, шт.'"),
]
NC = len(CALC)
fields = "".join(cache_field(i) for i in range(NF)) + "".join(f'<cacheField name="{escape(n_)}" numFmtId="0" formula="{escape(f_)}" databaseField="0"/>' for n_, f_ in CALC)
cache = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" refreshOnLoad="1" refreshedVersion="8" minRefreshableVersion="3" createdVersion="6" recordCount="{n}"><cacheSource type="worksheet"><worksheetSource ref="A1:{get_column_letter(NF)}{last}" sheet="Данные"/></cacheSource><cacheFields count="{NF + NC}">{fields}</cacheFields></pivotCacheDefinition>'
def pf_axis(i, ax): return f'<pivotField axis="{ax}" showAll="0"><items count="{len(sv[i])+1}">' + "".join(f'<item x="{k}"/>' for k in range(len(sv[i]))) + '<item t="default"/></items></pivotField>'
def table(name, data):
    dset = {f for _, f, _ in data}
    pf = []
    for i in range(NF + NC):
        if i in AX_ROW: pf.append(pf_axis(i, "axisRow"))
        elif i in AX_PAGE: pf.append(pf_axis(i, "axisPage"))
        elif i in dset: pf.append('<pivotField dataField="1" showAll="0"/>' if i < NF else '<pivotField dataField="1" dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        elif i >= NF: pf.append('<pivotField dragToRow="0" dragToCol="0" dragToPage="0" showAll="0" defaultSubtotal="0"/>')
        else: pf.append('<pivotField showAll="0"/>')
    dfs = "".join(f'<dataField name="{escape(nm)}" fld="{f}" baseField="0" baseItem="0" numFmtId="{nf}"/>' for nm, f, nf in data)
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="{name}" cacheId="1" showError="1" errorCaption="—" applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="Значения" updatedVersion="8" minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0" outline="1" outlineData="1" multipleFieldFilters="0" rowHeaderCaption="Вид / источник / аукцион / лот"><location ref="A4:V12" firstHeaderRow="0" firstDataRow="1" firstDataCol="1" rowPageCount="1" colPageCount="1"/><pivotFields count="{NF+NC}">{"".join(pf)}</pivotFields><rowFields count="4">' + "".join(f'<field x="{i}"/>' for i in AX_ROW) + f'</rowFields><colFields count="1"><field x="-2"/></colFields><pageFields count="1"><pageField fld="2" hier="-1"/></pageFields><dataFields count="{len(data)}">{dfs}</dataFields><pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/></pivotTableDefinition>'
c0 = NF
common_usd = [("Куплено с $, шт.", 16, 3), ("Куплено, $", 17, 3), ("Ср. молоток закупки, $ (по лотам с $)", c0 + 6, 3),
              ("Использовано с $, шт.", 18, 3), ("Использовано, $", 19, 3), ("Ср. молоток использования, $ (по лотам с $)", c0 + 7, 3),
              ("Остаток с $, шт.", 20, 3), ("Остаток, $", 21, 3), ("Ср. молоток остатка, $ (по лотам с $)", c0 + 8, 3)]
t1 = table("Сырьё с НДС", [("Куплено, шт.", 6, 3), ("Куплено, ₽", 7, 3), ("Ср. цена закупки, ₽/шт.", c0 + 0, 3), ("Использовано, шт.", 8, 3), ("Использовано, ₽", 9, 3), ("Ср. цена использования, ₽/шт.", c0 + 1, 3),
                           ("Остаток, шт.", 10, 3), ("Остаток, ₽", 11, 3), ("Ср. цена остатка, ₽/шт.", c0 + 2, 3)] + common_usd)
t2 = table("Сырьё без НДС", [("Куплено, шт.", 6, 3), ("Куплено, ₽", 12, 3), ("Ср. цена закупки, ₽/шт.", c0 + 3, 3), ("Использовано, шт.", 8, 3), ("Использовано, ₽", 13, 3), ("Ср. цена использования, ₽/шт.", c0 + 4, 3),
                           ("Остаток, шт.", 10, 3), ("Остаток, ₽", 14, 3), ("Ср. цена остатка, ₽/шт.", c0 + 5, 3)] + common_usd)
zi = zipfile.ZipFile(tmp); zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
for nm in zi.namelist():
    b = zi.read(nm)
    if nm == "[Content_Types].xml":
        t = b.decode(); add = '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/><Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/><Override PartName="/xl/pivotTables/pivotTable2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        b = t.replace("</Types>", add + "</Types>").encode()
    elif nm == "xl/workbook.xml":
        t = b.decode(); b = t.replace('</workbook>', '<pivotCaches><pivotCache xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" cacheId="1" r:id="rIdPC1"/></pivotCaches></workbook>').encode()
    elif nm == "xl/_rels/workbook.xml.rels":
        t = b.decode(); b = t.replace("</Relationships>", '<Relationship Id="rIdPC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode()
    zo.writestr(nm, b)
zo.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
zo.writestr("xl/pivotTables/pivotTable1.xml", t1); zo.writestr("xl/pivotTables/pivotTable2.xml", t2)
rel = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'
zo.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels", rel); zo.writestr("xl/pivotTables/_rels/pivotTable2.xml.rels", rel)
for idx, pt in ((1, 1), (2, 2)):
    zo.writestr(f"xl/worksheets/_rels/sheet{idx}.xml.rels", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" Target="../pivotTables/pivotTable{pt}.xml"/></Relationships>')
zo.close(); print("ok", n)
